import logging
import pandas as pd
import psycopg2
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
import time

from config import DB_CONFIG, FILE_PATHS

logger = logging.getLogger(__name__)

def clean_value(val):
    """通用的空值处理"""
    if pd.isna(val) or val is None:
        return ''
    return str(val).strip()

def read_and_update_excel(EXCEL_FILE, SHEET_NAME):
    logger.info("正在读取 Excel 文件...")
    df = None
    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_NAME, header=1, dtype=str)
        current_cols = df.columns.tolist()
        logger.info(f"读取到的列名: {current_cols[:10]}...")

    except Exception as e:
        logger.error(f"读取 Excel 失败: {e}")
        return

    if '发货ID' not in df.columns:
        logger.error("未找到 '发货ID' 列")
        return

    shipment_ids = df['发货ID'].dropna().tolist()
    if not shipment_ids:
        logger.warning("没有找到发货ID")
        return

    logger.info(f"找到 {len(shipment_ids)} 个发货ID，正在查询数据库...")

    # 2. 查询数据库
    db_results = {}
    try:
        conn = psycopg2.connect(**DB_CONFIG)
        cursor = conn.cursor()

        placeholders = ','.join(['%s'] * len(shipment_ids))
        query = f"""
            SELECT shipment_id, status, latest_info, vessel_voyage,
                   etd, eta, signed_at, is_inspected
            FROM logistics_shipments
            WHERE shipment_id IN ({placeholders})
        """
        cursor.execute(query, shipment_ids)
        rows = cursor.fetchall()

        for row in rows:
            db_results[row[0]] = {
                'status': clean_value(row[1]),
                'latest_info': clean_value(row[2]),
                'vessel_voyage': clean_value(row[3]),
                'etd': clean_value(row[4]),
                'eta': clean_value(row[5]),
                'signed_at': clean_value(row[6]),
                'is_inspected': bool(row[7])
            }
        conn.close()
        logger.info(f"数据库查询完成，找到 {len(db_results)} 条匹配记录")
    except Exception as e:
        logger.error(f"数据库查询失败: {e}")
        return

    field_to_df_col = {
        'status': '货件在途状态',
        'latest_info': '货件状态备注',
        'vessel_voyage': '船名航次',
        'etd': '开船/起飞日期ETD',
        'eta': '到港时间\n（ETA）',
        'signed_at': '妥投\n日期',
        'is_inspected': '是否有被查验'
    }

    # 3. 更新 DataFrame 数据
    for idx, row in df.iterrows():
        shipment_id = row['发货ID']
        if shipment_id in db_results:
            db_data = db_results[shipment_id]
            for db_field, df_col in field_to_df_col.items():
                if df_col in df.columns:
                    if db_field == 'is_inspected':
                        df.at[idx, df_col] = '是' if db_data[db_field] else '否'
                    else:
                        df.at[idx, df_col] = db_data[db_field]


    # 4. 写回 Excel
    output_file = EXCEL_FILE

    try:
        wb = load_workbook(EXCEL_FILE)
        ws = wb[SHEET_NAME]

        col_letter_mapping = {
            'AY': '开船/起飞日期ETD',
            'AZ': '妥投\n日期',
            'BC': '货件在途状态',
            'BD': '货件状态备注',
            'BE': '到港时间\n（ETA）',
            'BL': '船名航次',
            'BG': '是否有被查验'
        }

        start_excel_row = 3
        for i in range(len(shipment_ids)):
            excel_row = start_excel_row + i
            shipment_id = shipment_ids[i]

            if shipment_id in db_results:
                db_data = db_results[shipment_id]

                for col_letter, df_col_name in col_letter_mapping.items():
                    col_idx = column_index_from_string(col_letter)

                    if df_col_name in field_to_df_col.values():
                        db_field = [k for k, v in field_to_df_col.items() if v == df_col_name][0]
                        value = db_data[db_field]

                        if db_field == 'is_inspected':
                            value = '是' if value else ''

                        ws.cell(row=excel_row, column=col_idx, value=value)

        wb.save(output_file)
        logger.info(f"数据回填完成！文件已保存为: {output_file}")

    except Exception as e:
        logger.error(f"写入 Excel 失败: {e}")

if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
    EXCEL_FILE = FILE_PATHS['main_excel']
    SHEET_NAME = '发货数据详情'
    read_and_update_excel(EXCEL_FILE, SHEET_NAME)
